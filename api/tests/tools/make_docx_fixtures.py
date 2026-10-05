"""Regenerate the checked-in `.docx` fixtures under tests/fixtures/.

They are committed, so pytest needs nothing extra. Only re-run this if you want to change them:

    cd api && uv run --with python-docx python tests/tools/make_docx_fixtures.py

`python-docx` is deliberately *not* a project dependency — it only writes fixtures, the app only
reads them with mammoth.
"""

from __future__ import annotations

import struct
import sys
import zlib
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.run import Run

FIXTURES = Path(__file__).resolve().parent.parent / "fixtures"

HYPERLINK_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)


def tiny_png(width: int = 6, height: int = 4) -> bytes:
    """Smallest thing a .docx reader will accept as an image."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    scanlines = b"".join(
        b"\x00" + bytes((x * 37) % 256 for x in range(3 * width)) for _ in range(height)
    )
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", header)
        + chunk(b"IDAT", zlib.compress(scanlines))
        + chunk(b"IEND", b"")
    )


def add_link(paragraph, text: str, url: str) -> None:
    """python-docx has no link helper; a hyperlink is a run wrapped in <w:hyperlink>."""
    rel_id = paragraph.part.relate_to(url, HYPERLINK_REL, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = Run(OxmlElement("w:r"), paragraph)
    run.text = text
    link.append(run._r)
    paragraph._p.append(link)


def write(name: str, document: Document) -> None:
    buffer = BytesIO()
    document.save(buffer)
    path = FIXTURES / name
    path.write_bytes(buffer.getvalue())
    print(f"wrote {path.relative_to(FIXTURES.parent.parent)} ({len(buffer.getvalue())} bytes)")


def make_writeup() -> None:
    """A plausible staff write-up: every element the pipeline must survive."""
    doc = Document()
    doc.add_heading("Cleaning Up Watershed Sensor Exports Before Class", level=1)
    doc.add_paragraph(
        "The oddson-meter export is one 40,000-row CSV with a fresh header every time the "
        "logger reconnects. This script tidies it before students see it."
    )

    doc.add_heading("What you need", level=2)
    para = doc.add_paragraph()
    para.add_run("A Colby account with access to the ")
    para.add_run("sensor share").bold = True
    para.add_run(" and ")
    para.add_run("Python 3.12 or newer").italic = True
    para.add_run(".")

    doc.add_heading("Steps", level=2)
    for step in (
        "Download the export into ~/sensors/inbox.",
        "Run python clean_sensors.py inbox/export.csv.",
        "Paste the summary block into the LMS page.",
    ):
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("Checklist", level=2)
    for step in ("Check the headers.", "Check the timestamps.", "Flag sensor 4."):
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("Files", level=2)
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, (left, right) in enumerate(
        (
            ("File", "Purpose"),
            ("clean_sensors.py", "the script itself"),
            ("sensors.csv", "sensor id to name mapping"),
        )
    ):
        table.rows[row].cells[0].text = left
        table.rows[row].cells[1].text = right

    doc.add_heading("Where the docs live", level=2)
    para = doc.add_paragraph("The logger manual is on ")
    add_link(para, "the Colby field-equipment page", "https://www.colby.edu/field-equipment")
    para.add_run(".")

    doc.add_heading("Screenshot", level=2)
    doc.add_paragraph("This is what a bad export looks like:")
    doc.add_picture(BytesIO(tiny_png()), width=Inches(1))

    doc.add_heading("Gotchas", level=3)
    doc.add_paragraph(
        "Anything suspicious is only text, never markup: <script>alert(1)</script> and "
        "<img src=x onerror=alert(2)> must survive as words. So must a bare link written as "
        "<https://www.colby.edu/directory>."
    )
    write("writeup.docx", doc)


def make_empty() -> None:
    """Headings and body all missing -> the pipeline must call this unreadable."""
    doc = Document()
    doc.add_paragraph("")
    write("empty.docx", doc)


def make_plain_text() -> None:
    """Not a .docx at all: the bytes FastAPI must refuse before touching mammoth."""
    path = FIXTURES / "not-a-docx.docx"
    path.write_text(
        "This is a .txt file wearing a .docx name, which is exactly the case to reject.\n"
    )
    print(f"wrote {path.relative_to(FIXTURES.parent.parent)}")


def main() -> int:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    make_writeup()
    make_empty()
    make_plain_text()
    return 0


if __name__ == "__main__":
    sys.exit(main())
